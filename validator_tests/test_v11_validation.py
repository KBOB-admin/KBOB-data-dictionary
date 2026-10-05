import shutil
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts' / 'validator'))

from validate_strukturvorlage import Validator  # noqa: E402


class V11ValidationTests(unittest.TestCase):
    SOURCE = ROOT / 'templates' / '2026_09_Strukturvorlage_Data_Dictionary_leer_v1.1.0.xlsx'

    def make_workbook(self, override=None, enumeration_designation='Door states', base_type='STRING', ifc_type='IfcLabel'):
        directory = tempfile.TemporaryDirectory()
        path = Path(directory.name) / self.SOURCE.name
        shutil.copy2(self.SOURCE, path)
        workbook = load_workbook(path)

        properties = workbook['Properties']
        property_headers = {cell.value: cell.column for cell in properties[1] if cell.value}
        property_values = {
            'Designation (EN)': 'Door state',
            'Bezeichnung (DE)': 'Türstatus',
            'Description (EN)': 'State of the door',
            'Beschreibung (DE)': 'Status der Tür',
            'DataType\n(Base Type)': base_type,
            'DataType\n(IFC)': ifc_type,
            'EnumerationDesignation (EN)': enumeration_designation,
        }
        for header, value in property_values.items():
            properties.cell(7, property_headers[header]).value = value

        values = workbook['Values']
        value_headers = {cell.value: cell.column for cell in values[1] if cell.value}
        value_values = {
            'Designation (EN)': 'Door states',
            'Bezeichnung (DE)': 'Türstatuswerte',
            'Enumeration (EN)': '["Open", "Closed", "Locked"]',
            'Werteliste\n(DE)': '["Offen", "Geschlossen", "Verriegelt"]',
            'Provenance (PROV)': 'test',
        }
        for header, value in value_values.items():
            values.cell(7, value_headers[header]).value = value

        if override is not None:
            matrix = workbook['Data_Template']
            matrix.cell(2, 6).value = 'Door state'
            data_row = 6 if matrix.cell(3, 1).value == 'Validierung' else 5
            matrix.cell(data_row, 6).value = override

        workbook.save(path)
        workbook.close()
        return directory, path

    @staticmethod
    def codes(validator):
        return {finding.code for finding in validator.findings}

    def test_ifc_datatype_must_come_from_rules(self):
        directory, path = self.make_workbook(ifc_type='IfcWall')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_properties()
        self.assertIn('invalid_ifc_data_type', self.codes(validator))

    def test_valid_ifc43_value_datatype_is_accepted(self):
        directory, path = self.make_workbook(ifc_type='IfcBoolean')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_properties()
        self.assertNotIn('invalid_ifc_data_type', self.codes(validator))

    def test_property_resolves_enumeration_by_english_designation(self):
        directory, path = self.make_workbook(override='["Open"; "Locked"]')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_matrix()
        self.assertNotIn('override_without_global_enumeration', self.codes(validator))
        self.assertNotIn('invalid_allowed_values_override', self.codes(validator))

    def test_out_of_range_override_is_rejected(self):
        directory, path = self.make_workbook(override='["Open"; "Broken"]')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_matrix()
        self.assertIn('invalid_allowed_values_override', self.codes(validator))

    def test_localized_override_is_accepted_from_any_language_list(self):
        directory, path = self.make_workbook(override='["Offen"; "Verriegelt"]')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_matrix()
        self.assertNotIn('invalid_allowed_values_override', self.codes(validator))

    def test_override_without_connected_enumeration_is_rejected(self):
        directory, path = self.make_workbook(override='"Open"', enumeration_designation=None)
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_matrix()
        self.assertIn('override_without_global_enumeration', self.codes(validator))

    def test_boolean_override_is_inferred_without_explicit_enumeration(self):
        directory, path = self.make_workbook(
            override='["true"]',
            enumeration_designation=None,
            base_type='BOOLEAN',
            ifc_type='IfcBoolean',
        )
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_matrix()
        codes = self.codes(validator)
        self.assertNotIn('override_without_global_enumeration', codes)
        self.assertNotIn('invalid_allowed_values_override', codes)
        self.assertIn('boolean_enumeration_inferred_from_datatype', codes)

    def test_boolean_base_type_rejects_ifc_logical(self):
        directory, path = self.make_workbook(base_type='BOOLEAN', ifc_type='IfcLogical')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_properties()
        self.assertIn('incompatible_property_data_types', self.codes(validator))

    def test_governance_findings_are_emitted_once_per_populated_row(self):
        directory, path = self.make_workbook(override='["Open"; "Closed"]')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_matrix()
        codes = [finding.code for finding in validator.findings]
        self.assertEqual(1, codes.count('matrix_missing_status'))
        self.assertEqual(1, codes.count('matrix_missing_version_date'))
        self.assertEqual(1, codes.count('matrix_missing_provenance'))

    def test_unknown_enumeration_designation_is_rejected(self):
        directory, path = self.make_workbook(enumeration_designation='Missing states')
        self.addCleanup(directory.cleanup)
        validator = Validator(path)
        validator.validate_properties()
        self.assertIn('unknown_enumeration_designation', self.codes(validator))


if __name__ == '__main__':
    unittest.main()
