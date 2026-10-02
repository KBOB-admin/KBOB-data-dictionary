import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts' / 'validator'))

from validate_strukturvorlage import USER_DEFINED_IFC_URI, Validator  # noqa: E402


class IfcMappingTests(unittest.TestCase):
    SOURCE = ROOT / 'templates' / '2026_09_Strukturvorlage_Data_Dictionary_leer_v1.1.0.xlsx'
    USE_CASE_SOURCE = ROOT / 'WIP data dictionaries' / 'IFMA' / 'Use Case Grundlagen Ausschreibung H-K_v0.6.xlsx'

    def setUp(self):
        self.validator = Validator(self.SOURCE)

    def finding_codes(self):
        return {finding.code for finding in self.validator.findings}

    def test_v11_empty_formula_property_headings_are_ignored(self):
        self.validator.validate_matrix()
        self.assertNotIn('matrix_unknown_property_label', self.finding_codes())

    def test_v11_datatemplate_guidance_and_metadata_structure(self):
        sheet = self.validator.wb['Data_Template']
        self.assertEqual('Validierung', sheet['A3'].value)
        self.assertEqual(6, self.validator._sheet_start_row('Data_Template'))
        self.assertFalse(any(sheet.cell(2, column).value == 0 for column in range(6, 53)))
        structure = self.validator.parse_datatemplate_structure()
        self.assertEqual(77, structure['document_block_anchor'])
        self.assertEqual(83, structure['loin_anchor'])
        self.assertEqual(88, structure['governance_anchor'])

    def test_ifc43_value_datatype_rules_are_bound_and_complete(self):
        values = self.validator._load_dropdown_values('IFC Data Type')
        self.assertIn('IfcBoolean', values)
        self.assertIn('IfcLabel', values)
        self.assertIn('IfcLengthMeasure', values)
        self.assertIn('nicht definiert', values)
        self.assertNotIn('IfcWall', values)
        self.assertNotIn('IfcWindowTypeEnum', values)
        self.assertEqual(110, len(values))

    def test_semicolon_override_accepts_quoted_strings(self):
        mode, values = self.validator.parse_datatemplate_override('["Wert 1"; "Wert 2"]', 'Data_Template', 5, 'col-6')
        self.assertEqual(('subset', ['Wert 1', 'Wert 2']), (mode, values))
        self.assertEqual(set(), self.finding_codes())

    def test_semicolon_override_accepts_unquoted_numbers(self):
        mode, values = self.validator.parse_datatemplate_override('[1.2; 2.2; 2.3]', 'Data_Template', 5, 'col-6')
        self.assertEqual(('subset', ['1.2', '2.2', '2.3']), (mode, values))
        self.assertEqual(set(), self.finding_codes())

    def test_comma_override_is_accepted_with_formatting_notice(self):
        mode, values = self.validator.parse_datatemplate_override('[1.2, 2.2]', 'Data_Template', 5, 'col-6')
        self.assertEqual(('subset', ['1.2', '2.2']), (mode, values))
        self.assertIn('noncanonical_allowed_values_separator', self.finding_codes())

    def test_empty_override_means_complete_enumeration_with_notice(self):
        mode, values = self.validator.parse_datatemplate_override('[]', 'Data_Template', 5, 'col-6')
        self.assertEqual(('all', []), (mode, values))
        self.assertIn('empty_allowed_values_override', self.finding_codes())

    def test_x_is_assignment_without_local_override(self):
        mode, values = self.validator.parse_datatemplate_override('X', 'Data_Template', 5, 'col-6')
        self.assertEqual(('assignment', []), (mode, values))
        self.assertEqual(set(), self.finding_codes())

    def test_malformed_override_is_rejected_separately(self):
        mode, values = self.validator.parse_datatemplate_override('[Wert 1; "Wert 2"]', 'Data_Template', 5, 'col-6')
        self.assertEqual(('invalid', []), (mode, values))
        self.assertIn('malformed_allowed_values_override', self.finding_codes())

    def test_type_object_entity_column_is_not_validated(self):
        validator = Validator(self.USE_CASE_SOURCE)
        report = validator.validate()
        finding_codes = {finding['code'] for finding in report['findings']}
        self.assertNotIn('invalid_ifc_type_object_entity', finding_codes)
        self.assertNotIn('invalid_ifc_object_type_pair', finding_codes)
        self.assertNotIn('missing_ifc_object_entity', finding_codes)
        self.assertNotIn('minimum_documents', finding_codes)

    def test_v06_predefined_types_are_covered_by_ifc_uri_cache(self):
        validator = Validator(self.USE_CASE_SOURCE)
        report = validator.validate()
        finding_codes = {finding['code'] for finding in report['findings']}
        self.assertNotIn('unknown_ifc_uri', finding_codes)
        self.assertNotIn('invalid_predefined_type', finding_codes)

    def test_user_defined_is_an_accepted_ifc_uri_value(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workbook = openpyxl.load_workbook(self.USE_CASE_SOURCE)
            ws = workbook['Classes']
            headers = self.validator._sheet_headers('Classes')
            ifc_uri_column = headers['IFC_URI']
            for row_idx in range(self.validator._sheet_start_row('Classes'), ws.max_row + 1):
                if ws.cell(row_idx, ifc_uri_column).value in (None, ''):
                    ws.cell(row_idx, ifc_uri_column).value = USER_DEFINED_IFC_URI
            test_path = Path(temp_dir) / 'user-defined-ifc-uri.xlsx'
            workbook.save(test_path)

            report = Validator(test_path).validate()
            finding_codes = {finding['code'] for finding in report['findings']}
            self.assertNotIn('missing_ifc_uri', finding_codes)
            self.assertNotIn('invalid_ifc_uri', finding_codes)
            self.assertNotIn('invalid_ifc_uri_namespace', finding_codes)

    def test_predefined_type_is_validated_separately_from_type_entity(self):
        uri = 'https://identifier.buildingsmart.org/uri/buildingsmart/ifc/4.3/class/IfcTankVESSEL'
        self.validator.validate_predefined_type('IfcTank', 'VESSEL', uri, 'Classes', 7)
        self.assertEqual(set(), self.finding_codes())

    def test_invalid_predefined_type_is_rejected(self):
        uri = 'https://identifier.buildingsmart.org/uri/buildingsmart/ifc/4.3/class/IfcTankNOTAREALTYPE'
        self.validator.validate_predefined_type('IfcTank', 'NOTAREALTYPE', uri, 'Classes', 7)
        self.assertIn('invalid_predefined_type', self.finding_codes())

    def test_single_ifc_set_reference_is_accepted(self):
        references = self.validator.parse_ifc_linked_references('Pset_PipeSegmentTypeCommon', 'Properties', 60)
        self.assertEqual(['Pset_PipeSegmentTypeCommon'], references)
        self.assertEqual(set(), self.finding_codes())

    def test_multiple_ifc_set_references_require_json_array(self):
        references = self.validator.parse_ifc_linked_references('Pset_PipeSegmentTypeCommon\nQto_WallBaseQuantities', 'Properties', 60)
        self.assertEqual([], references)
        self.assertIn('invalid_ifc_linked_list_syntax', self.finding_codes())

    def test_json_ifc_set_reference_list_is_accepted(self):
        raw = '["Pset_PipeSegmentTypeCommon", "Qto_WallBaseQuantities"]'
        references = self.validator.parse_ifc_linked_references(raw, 'Properties', 60)
        self.assertEqual(['Pset_PipeSegmentTypeCommon', 'Qto_WallBaseQuantities'], references)
        self.assertEqual(set(), self.finding_codes())

    def test_ifc_set_reference_list_requires_non_empty_strings(self):
        references = self.validator.parse_ifc_linked_references('["Pset_WallCommon", ""]', 'Properties', 60)
        self.assertEqual([], references)
        self.assertIn('invalid_ifc_linked_list_syntax', self.finding_codes())

    def test_duplicate_ifc_set_reference_is_rejected(self):
        raw = '["Qto_WallBaseQuantities", "Qto_WallBaseQuantities"]'
        references = self.validator.parse_ifc_linked_references(raw, 'Properties', 60)
        self.assertEqual(['Qto_WallBaseQuantities', 'Qto_WallBaseQuantities'], references)
        self.assertIn('duplicate_ifc_linked_reference', self.finding_codes())


if __name__ == '__main__':
    unittest.main()
